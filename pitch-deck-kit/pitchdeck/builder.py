"""Build navy/gold investment-pitch decks from a YAML/JSON content spec.

Structure of the generated .pptx
--------------------------------
* Theme   - brand colors and Garamond written into theme1.xml.
* Master  - target logo (top right) and sponsor wordmark (footer) live on the
            master, so they cannot be nudged out of place slide by slide.
* Layouts - "Cover", "Divider", "Content", "Blank". Titles are real title
            placeholders (they show in Outline view and to screen readers).
* Slides  - section tracker, page number and all body components are drawn
            per slide from the spec.

Entry point: ``build_deck(spec_dict, base_dir, out_path) -> list[str]`` returns
warnings (text that probably overflows, missing images, unknown sections).
"""
from __future__ import annotations

import math
import os
import re
from dataclasses import replace

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.shapes.autoshape import CT_Shape
from pptx.oxml.shapes.connector import CT_Connector
from pptx.oxml.shapes.picture import CT_Picture
from pptx.shapes.autoshape import Shape
from pptx.shapes.connector import Connector
from pptx.util import Emu, Inches, Pt

from .style import CHART_SERIES, GEO, MIN_BODY_SIZE, MIN_TABLE_SIZE, PALETTE, THEME_COLORS, TYPE, Brand, TypeStyle

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
ALIGN = {"l": PP_ALIGN.LEFT, "left": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "center": PP_ALIGN.CENTER,
         "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "right": PP_ALIGN.RIGHT}
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
MARKUP = re.compile(r"(\*\*.+?\*\*|__.+?__)")


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


def ink_on(hex_: str) -> str:
    """Readable label color (white or title gray) for text sitting on hex_."""
    r, g, b = (int(hex_[i:i + 2], 16) / 255 for i in (0, 2, 4))
    lum = 0.2126 * r + 0.7152 * g + 0.0722 * b
    return PALETTE["white"] if lum < 0.5 else PALETTE["title"]


def I(v: float) -> Emu:  # noqa: E743 - inches helper, used everywhere
    return Inches(v)


# ==========================================================================
# Text helpers
# ==========================================================================
def parse_markup(text: str):
    """'a **bold** and __underlined__ word' -> [(text, bold, underline), ...]."""
    out = []
    for tok in MARKUP.split(str(text)):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            out.append((tok[2:-2], True, False))
        elif tok.startswith("__") and tok.endswith("__"):
            out.append((tok[2:-2], True, True))
        else:
            out.append((tok, False, False))
    return out


def plain(text: str) -> str:
    return "".join(t for t, _, _ in parse_markup(text))


def est_lines(text: str, width_in: float, size_pt: float, small_caps: bool = False) -> int:
    """Greedy word-wrap estimate using an average glyph width."""
    em = GEO.em_small_caps if small_caps else GEO.em_regular
    char_w = size_pt * em / 72.0
    cap = max(1, int(width_in / char_w))
    lines, cur = 0, 0
    for para in plain(text).split("\n"):
        lines += 1
        cur = 0
        for word in para.split():
            need = len(word) + (1 if cur else 0)
            if cur + need > cap:
                lines += 1
                cur = len(word)
            else:
                cur += need
    return max(lines, 1)


def line_h(size_pt: float, spacing: float = 1.0) -> float:
    return size_pt * GEO.line_height * spacing / 72.0


def style_run(run, ts: TypeStyle, font: str, bold=None, underline=None, color=None, size=None):
    f = run.font
    f.name = font
    f.size = Pt(size or ts.size)
    f.bold = ts.bold if bold is None else bold
    f.underline = ts.underline if underline is None else underline
    f.color.rgb = rgb(color or ts.color)
    rPr = run._r.get_or_add_rPr()
    if ts.small_caps:
        rPr.set("cap", "small")
    for tag in ("a:ea", "a:cs"):
        if rPr.find(qn(tag)) is None:
            el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", font)


def add_runs(paragraph, text, ts: TypeStyle, font, color=None, size=None):
    for chunk, b, u in parse_markup(text):
        r = paragraph.add_run()
        r.text = chunk
        style_run(r, ts, font, bold=(ts.bold or b), underline=(ts.underline or u), color=color, size=size)


def set_bullet(paragraph, level: int, color: str, size: float, font: str):
    pPr = paragraph._p.get_or_add_pPr()
    marL = GEO.bullet_marL if level == 0 else GEO.bullet_l2_marL
    pPr.set("marL", str(int(I(marL))))
    pPr.set("indent", str(-int(I(GEO.bullet_hang))))
    for tag in ("a:buClr", "a:buSzPts", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buSzPts")).set("val", str(int(size * 100)))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", font)
    etree.SubElement(pPr, qn("a:buChar")).set("char", "●" if level == 0 else "○")


def no_bullet(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    if pPr.find(qn("a:buNone")) is None:
        etree.SubElement(pPr, qn("a:buNone"))


def frame(tf, inset_lr=GEO.inset, inset_tb=GEO.inset, anchor=MSO_ANCHOR.TOP, wrap=True):
    tf.margin_left = tf.margin_right = I(inset_lr)
    tf.margin_top = tf.margin_bottom = I(inset_tb)
    tf.vertical_anchor = anchor
    tf.word_wrap = wrap
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for child in list(bodyPr):
        bodyPr.remove(child)
    etree.SubElement(bodyPr, qn("a:noAutofit"))


def fill_shape(shape, fill: str | None, line: str | None = None, line_pt: float = 1.0, dash=None):
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_pt)
        if dash:
            shape.line.dash_style = dash
    else:
        shape.line.fill.background()
    # kill the theme's default shadow / style reference
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


# ==========================================================================
# Builder
# ==========================================================================
class DeckBuilder:
    def __init__(self, spec: dict, base_dir: str):
        self.spec = spec
        self.base = base_dir
        b = spec.get("brand", {})
        self.brand = Brand(
            font=b.get("font", "Garamond"),
            number_font=b.get("number_font", "Calibri"),
            target_name=b.get("target_name", ""),
            target_logo=self._path(b.get("target_logo")),
            sponsor_name=b.get("sponsor_name", ""),
            sponsor_logo=self._path(b.get("sponsor_logo")),
        )
        self.sections = spec.get("sections", [])
        self.warnings: list[str] = []
        self.prs = Presentation()
        self.prs.slide_width = I(GEO.slide_w)
        self.prs.slide_height = I(GEO.slide_h)
        self.divider_count = 0
        self.slide_no = 0
        self._patch_theme()
        self._build_master_and_layouts()

    # ---------------------------------------------------------------- utils
    def _path(self, p):
        if not p:
            return None
        full = p if os.path.isabs(p) else os.path.join(self.base, p)
        return full

    def warn(self, msg):
        self.warnings.append(f"slide {self.slide_no}: {msg}")

    @property
    def font(self):
        return self.brand.font

    # ---------------------------------------------------------------- theme
    def _patch_theme(self):
        part = self.prs.slide_master.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        ns = {"a": A_NS}
        root.set("name", "Pitch Navy")
        scheme = root.find(".//a:clrScheme", ns)
        scheme.set("name", "Pitch Navy")
        for slot, hex_ in THEME_COLORS.items():
            el = scheme.find(f"a:{slot}", ns)
            for c in list(el):
                el.remove(c)
            etree.SubElement(el, f"{{{A_NS}}}srgbClr").set("val", hex_)
        fs = root.find(".//a:fontScheme", ns)
        fs.set("name", "Pitch Navy")
        for which in ("majorFont", "minorFont"):
            fs.find(f"a:{which}/a:latin", ns).set("typeface", self.font)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ------------------------------------------------- master-level drawing
    def _add_pic(self, owner, path, x, y, w, h, name):
        _, rId = owner.part.get_or_add_image_part(path)
        sid = owner.shapes._next_shape_id
        pic = CT_Picture.new_pic(sid, name, name, rId, I(x), I(y), I(w), I(h))
        owner.shapes._spTree.append(pic)
        return pic

    def _add_rect(self, owner, x, y, w, h, name, fill, line=None, line_pt=1.0, prst="rect"):
        sid = owner.shapes._next_shape_id
        sp = CT_Shape.new_autoshape_sp(sid, name, prst, I(x), I(y), I(w), I(h))
        owner.shapes._spTree.append(sp)
        shp = Shape(sp, owner.shapes)
        fill_shape(shp, fill, line, line_pt)
        return shp

    def _add_line(self, owner, x, y, w, name, color, pt):
        sid = owner.shapes._next_shape_id
        cx = CT_Connector.new_cxnSp(sid, name, "line", I(x), I(y), I(w), 0, False, False)
        owner.shapes._spTree.append(cx)
        c = Connector(cx, owner.shapes)
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(pt)
        return c

    def _logo_size(self, path, h, max_w):
        with Image.open(path) as im:
            w = h * im.width / im.height
        if w > max_w:
            h, w = h * max_w / w, max_w
        return w, h

    def _set_style_lvl(self, parent, tag, ts: TypeStyle, algn="l", anchor_extra=None, marL=0, indent=0,
                       bullet=False, spacing=100):
        """Replace <a:lvl1pPr> (or given tag) inside a lstStyle/titleStyle."""
        old = parent.find(qn(tag))
        new = etree.Element(qn(tag))
        new.set("marL", str(int(I(marL))))
        new.set("indent", str(int(I(indent))))
        new.set("algn", algn)
        ln = etree.SubElement(new, qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPct")).set("val", str(spacing * 1000))
        sb = etree.SubElement(new, qn("a:spcBef"))
        etree.SubElement(sb, qn("a:spcPts")).set("val", "0")
        if bullet:
            c = etree.SubElement(new, qn("a:buClr"))
            etree.SubElement(c, qn("a:srgbClr")).set("val", ts.color)
            etree.SubElement(new, qn("a:buFont")).set("typeface", self.font)
            etree.SubElement(new, qn("a:buChar")).set("char", "●")
        else:
            etree.SubElement(new, qn("a:buNone"))
        d = etree.SubElement(new, qn("a:defRPr"))
        d.set("sz", str(int(ts.size * 100)))
        d.set("b", "1" if ts.bold else "0")
        d.set("cap", "small" if ts.small_caps else "none")
        sf = etree.SubElement(d, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", ts.color)
        for t in ("a:latin", "a:ea", "a:cs"):
            etree.SubElement(d, qn(t)).set("typeface", self.font)
        if old is not None:
            parent.replace(old, new)
        else:
            parent.insert(0, new)

    def _ph_lststyle(self, ph, ts, algn="l", anchor="b", marL=0, indent=0, insets=(GEO.inset, 0.05)):
        txBody = ph._element.find(qn("p:txBody"))
        bodyPr = txBody.find(qn("a:bodyPr"))
        bodyPr.set("anchor", anchor)
        bodyPr.set("lIns", str(int(I(insets[0]))))
        bodyPr.set("rIns", str(int(I(insets[0]))))
        bodyPr.set("tIns", str(int(I(insets[1]))))
        bodyPr.set("bIns", str(int(I(insets[1]))))
        bodyPr.set("wrap", "square")
        for c in list(bodyPr):
            bodyPr.remove(c)
        etree.SubElement(bodyPr, qn("a:noAutofit"))
        lst = txBody.find(qn("a:lstStyle"))
        if lst is None:
            lst = etree.Element(qn("a:lstStyle"))
            bodyPr.addnext(lst)
        self._set_style_lvl(lst, "a:lvl1pPr", ts, algn=algn, marL=marL, indent=indent)

    def _strip_placeholders(self, layout, keep_types):
        for ph in list(layout.placeholders):
            if ph.placeholder_format.type not in keep_types:
                ph._element.getparent().remove(ph._element)

    def _rename(self, layout, name):
        layout._element.find(qn("p:cSld")).set("name", name)

    def _build_master_and_layouts(self):
        from pptx.enum.shapes import PP_PLACEHOLDER as PH
        g, m = GEO, self.prs.slide_master
        # master text styles: titles + bulleted body
        txStyles = m._element.find(qn("p:txStyles"))
        self._set_style_lvl(txStyles.find(qn("p:titleStyle")), "a:lvl1pPr", TYPE["title"])
        self._set_style_lvl(txStyles.find(qn("p:bodyStyle")), "a:lvl1pPr", TYPE["body"],
                            marL=g.bullet_marL, indent=-g.bullet_hang, bullet=True)
        mt = m.placeholders[0]
        mt.left, mt.top, mt.width, mt.height = I(g.title_x), I(g.title_y), I(g.title_w), I(g.title_h)
        bodyPr = mt._element.find(qn("p:txBody")).find(qn("a:bodyPr"))
        bodyPr.set("anchor", "b")
        # brand marks on the master
        if self.brand.target_logo:
            w, h = self._logo_size(self.brand.target_logo, g.logo_h, g.logo_max_w)
            self._add_pic(m, self.brand.target_logo, g.slide_w - 0.307 - w, g.logo_y, w, h, "Target Logo")
        elif self.brand.target_name:
            self._master_text(m, self.brand.target_name, g.slide_w - 0.307 - 2.5, g.logo_y - 0.05, 2.5, 0.4,
                              TypeStyle(18, True, PALETTE["logo_navy"]), "r", "Target Name")
        if self.brand.sponsor_logo:
            w, h = self._logo_size(self.brand.sponsor_logo, g.sponsor_h, 2.2)
            self._add_pic(m, self.brand.sponsor_logo, g.sponsor_right - w, g.sponsor_y, w, h, "Sponsor Wordmark")
        elif self.brand.sponsor_name:
            self._master_text(m, self.brand.sponsor_name, g.sponsor_right - 2.5, g.page_y, 2.5, g.page_h,
                              TypeStyle(12, False, PALETTE["sponsor_text"], small_caps=True), "r", "Sponsor Name")

        layouts = {l.name: l for l in self.prs.slide_layouts}
        # Cover: no master marks, ticker in ctrTitle, byline in subTitle
        cover = layouts["Title Slide"]
        self._rename(cover, "Cover")
        cover._element.set("showMasterSp", "0")
        self._strip_placeholders(cover, {PH.CENTER_TITLE, PH.SUBTITLE})
        t, s = cover.placeholders[0], cover.placeholders[1]
        t.left, t.top, t.width, t.height = I(1.667), I(2.55), I(10.0), I(0.68)
        self._ph_lststyle(t, TYPE["cover_ticker"], algn="ctr", anchor="b")
        s.left, s.top, s.width, s.height = I(1.667), I(6.25), I(10.0), I(0.56)
        self._ph_lststyle(s, TYPE["byline"], algn="ctr", anchor="b")
        # Divider: navy band + gold stripe behind a white title
        div = layouts["Section Header"]
        self._rename(div, "Divider")
        self._strip_placeholders(div, {PH.TITLE})
        spTree = div.shapes._spTree
        band = self._add_rect(div, -0.042, g.band_y, g.slide_w + 0.084, g.band_h, "Divider Band",
                              PALETTE["primary"], PALETTE["divider_outline"], 0.75)
        stripe = self._add_rect(div, -0.042, g.band_y + g.band_h, g.slide_w + 0.084, g.stripe_h,
                                "Divider Stripe", PALETTE["gold"])
        for shp in (stripe, band):  # send behind the title placeholder
            spTree.remove(shp._element)
            spTree.insert(2, shp._element)
        t = div.placeholders[0]
        t.left, t.top, t.width, t.height = I(g.divider_text_x), I(g.band_y), I(g.divider_text_w), I(g.band_h)
        self._ph_lststyle(t, TYPE["divider"], anchor="ctr", marL=0.75, indent=-0.75, insets=(GEO.inset, GEO.inset))
        # Content: title placeholder + rule
        con = layouts["Title Only"]
        self._rename(con, "Content")
        self._strip_placeholders(con, {PH.TITLE})
        t = con.placeholders[0]
        t.left, t.top, t.width, t.height = I(g.title_x), I(g.title_y), I(g.title_w), I(g.title_h)
        self._ph_lststyle(t, TYPE["title"], anchor="b")
        self._add_line(con, g.rule_x, g.rule_y, g.rule_w, "Title Rule", PALETTE["ink"], g.rule_pt)
        blank = layouts["Blank"]
        self._strip_placeholders(blank, set())
        for name, l in layouts.items():
            if name not in ("Title Slide", "Section Header", "Title Only", "Blank"):
                self.prs.slide_layouts.remove(l)
        self.L = {"cover": cover, "divider": div, "content": con, "blank": blank}

    def _master_text(self, owner, text, x, y, w, h, ts, algn, name):
        shp = self._add_rect(owner, x, y, w, h, name, None)
        tf = shp.text_frame
        frame(tf, anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        p.alignment = ALIGN[algn]
        add_runs(p, text, ts, self.font)

    # ---------------------------------------------------- slide primitives
    def textbox(self, slide, x, y, w, h, name=None, anchor=MSO_ANCHOR.TOP, inset_tb=GEO.inset):
        tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
        if name:
            tb.name = name
        frame(tb.text_frame, anchor=anchor, inset_tb=inset_tb)
        return tb

    def rect(self, slide, x, y, w, h, fill, line=None, line_pt=1.0, shape=MSO_SHAPE.RECTANGLE, name=None):
        s = slide.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
        fill_shape(s, fill, line, line_pt)
        if name:
            s.name = name
        return s

    def label_rect(self, slide, x, y, w, h, text, ts: TypeStyle, fill, line=None, line_pt=1.0,
                   algn="ctr", shape=MSO_SHAPE.RECTANGLE, name=None, size=None, inset_tb=0.05, spacing=None):
        s = self.rect(slide, x, y, w, h, fill, line, line_pt, shape, name)
        tf = s.text_frame
        frame(tf, anchor=MSO_ANCHOR.MIDDLE, inset_tb=inset_tb)
        for i, line_text in enumerate(str(text).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = ALIGN[algn]
            if spacing:
                p.line_spacing = spacing
            no_bullet(p)
            add_runs(p, line_text, ts, self.font, size=size)
        return s

    def page_number(self, slide):
        g = GEO
        tb = self.textbox(slide, g.page_x, g.page_y, g.page_w, g.page_h, "Page Number",
                          anchor=MSO_ANCHOR.MIDDLE, inset_tb=0.05)
        tb.text_frame.word_wrap = False
        p = tb.text_frame.paragraphs[0]
        ts = TYPE["page_number"]
        r = p.add_run()
        r.text = "|  "
        style_run(r, ts, self.brand.number_font)
        fld = etree.SubElement(p._p, qn("a:fld"))
        fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
        fld.set("type", "slidenum")
        rPr = etree.SubElement(fld, qn("a:rPr"))
        rPr.set("lang", "en-US")
        rPr.set("sz", str(int(ts.size * 100)))
        sf = etree.SubElement(rPr, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", ts.color)
        etree.SubElement(rPr, qn("a:latin")).set("typeface", self.brand.number_font)
        etree.SubElement(fld, qn("a:t")).text = str(self.slide_no)

    def tracker(self, slide, active):
        if not self.sections:
            return
        if active and active not in self.sections:
            self.warn(f"section '{active}' is not in the deck's sections list")
        g, n = GEO, len(self.sections)
        w = (g.tracker_x1 - g.tracker_x0 - (n - 1) * g.tracker_gap) / n
        for i, label in enumerate(self.sections):
            x = g.tracker_x0 + i * (w + g.tracker_gap)
            on = label == active
            fill = PALETTE["active"] if on else PALETTE["neutral"]
            size = TYPE["tracker"].size
            if est_lines(label, w - 0.2, size) > 1:
                size = 12
            self.label_rect(slide, x, g.tracker_y, w, g.tracker_h, label, TYPE["tracker"], fill, fill,
                            name=f"Tracker: {label}{' (active)' if on else ''}", size=size)

    def notes(self, slide, text):
        if text:
            slide.notes_slide.notes_text_frame.text = str(text)

    def new_slide(self, layout_key):
        self.slide_no += 1
        s = self.prs.slides.add_slide(self.L[layout_key])
        return s

    def set_title(self, slide, text, max_w=GEO.title_w):
        ph = slide.shapes.title
        ts = TYPE["title"]
        size = ts.size
        while size > 24 and est_lines(text, max_w - 0.2, size, small_caps=True) > 1:
            size -= 2
        if est_lines(text, max_w - 0.2, size, small_caps=True) > 1:
            self.warn(f"title still wraps at {size}pt - shorten it: '{text}'")
        elif size != ts.size:
            self.warn(f"title shrunk to {size}pt to stay on one line: '{text}'")
        p = ph.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = str(text)
        if size != ts.size:
            r.font.size = Pt(size)

    # ======================================================== slide types
    def cover(self, s):
        g = GEO
        slide = self.new_slide("cover")
        if self.brand.target_logo:
            w, h = self._logo_size(self.brand.target_logo, 1.223, 5.98)
            slide.shapes.add_picture(self.brand.target_logo, I((g.slide_w - w) / 2), I(0.689), I(w), I(h))
        else:
            tb = self.textbox(slide, 1.667, 0.8, 10.0, 1.1, anchor=MSO_ANCHOR.MIDDLE)
            p = tb.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            add_runs(p, self.brand.target_name, TypeStyle(60, True, PALETTE["logo_navy"]), self.font)
        slide.shapes.title.text_frame.paragraphs[0].add_run().text = s.get("ticker", self.brand.target_name)
        kpis = s.get("kpis", [])
        if kpis:
            ts = TYPE["cover_kpi"]
            h = len(kpis) * line_h(ts.size, 1.5) + 2 * g.inset + 0.05
            longest = max(len(plain(k)) for k in kpis)
            w = min(10.0, max(4.656, longest * ts.size * g.em_small_caps / 72 + 0.6))
            box = self.rect(slide, (g.slide_w - w) / 2, 3.645, w, h, None, PALETTE["logo_navy"], 1.5, name="KPI Box")
            tf = box.text_frame
            frame(tf, anchor=MSO_ANCHOR.MIDDLE)
            for i, k in enumerate(kpis):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                p.line_spacing = 1.5
                no_bullet(p)
                add_runs(p, k, ts, self.font)
        byline = s.get("byline")
        if byline:
            slide.placeholders[1].text_frame.paragraphs[0].add_run().text = byline
        if self.brand.sponsor_logo:  # the cover layout hides master marks, so add the wordmark here
            w, h = self._logo_size(self.brand.sponsor_logo, g.sponsor_h, 2.2)
            slide.shapes.add_picture(self.brand.sponsor_logo, I(g.sponsor_right - w), I(g.sponsor_y), I(w), I(h))
        elif self.brand.sponsor_name:
            self._master_text(slide, self.brand.sponsor_name, g.sponsor_right - 2.5, g.page_y, 2.5, g.page_h,
                              TypeStyle(12, False, PALETTE["sponsor_text"], small_caps=True), "r", "Sponsor Name")
        if s.get("date"):
            tb = self.textbox(slide, 1.667, 6.82, 10.0, 0.3)
            p = tb.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            add_runs(p, s["date"], TypeStyle(14, False, PALETTE["title"]), self.font)
        self.page_number(slide)
        self.notes(slide, s.get("notes"))

    def agenda(self, s):
        g = GEO
        slide = self.new_slide("content")
        self.set_title(slide, s.get("title", "Agenda"))
        items = s.get("items") or [d.get("title") for d in self.spec["slides"] if d.get("type") == "divider"]
        n = len(items)
        top, bottom = 1.18, 6.35
        pitch = min(0.86, (bottom - top) / max(n - 1, 1)) if n > 1 else 0
        for i, item in enumerate(items):
            y = top + i * pitch
            self.label_rect(slide, 0.934, y, 0.584, 0.599, str(i + 1), TYPE["agenda_num"], PALETTE["ink"],
                            PALETTE["gold"], 1.5, shape=MSO_SHAPE.OVAL, name=f"Agenda {i+1}", inset_tb=0.0)
            tb = self.textbox(slide, 1.78, y + 0.03, 10.5, 0.54, anchor=MSO_ANCHOR.MIDDLE, inset_tb=0.0)
            p = tb.text_frame.paragraphs[0]
            no_bullet(p)
            add_runs(p, item, TYPE["agenda_item"], self.font)
        self.page_number(slide)
        self.notes(slide, s.get("notes"))

    def divider(self, s):
        slide = self.new_slide("divider")
        number = s.get("number") or ROMAN[self.divider_count]
        self.divider_count += 1
        slide.shapes.title.text_frame.paragraphs[0].add_run().text = f"{number}.\t{s['title']}"
        self.page_number(slide)
        self.notes(slide, s.get("notes"))

    def content(self, s):
        g = GEO
        slide = self.new_slide("content")
        self.set_title(slide, s["title"])
        has_tracker = s.get("tracker", True) and bool(self.sections) and s.get("section") is not None
        y1 = g.body_y1_tracker if has_tracker else g.body_y1_plain
        y0 = g.body_y0
        for b in s.get("top", []):  # full-width blocks above the columns
            full_w = g.body_x1 - g.body_x0
            h = self._height(b, full_w, TYPE["body"].size) or 1.0
            getattr(self, "_b_" + self._block_kind(b))(slide, g.body_x0, y0, full_w, h, b, TYPE["body"].size)
            y0 += h + g.block_gap
        if s.get("columns") and s.get("blocks"):
            self.warn("slide has both 'blocks' and 'columns'; 'blocks' was ignored (use 'top' for full-width blocks)")
        cols = s.get("columns") or [{"blocks": s.get("blocks", [])}]
        cols = [c if isinstance(c, dict) else {"blocks": c} for c in cols]
        weights = [c.get("width", 1) for c in cols]
        total_w = g.body_x1 - g.body_x0 - g.gutter * (len(cols) - 1)
        x = g.body_x0
        plans = []
        for c, wt in zip(cols, weights):
            w = total_w * wt / sum(weights)
            plans.append((x, w, c.get("blocks", [])))
            x += w + g.gutter
        # one body size for the whole slide so columns stay consistent
        size = s.get("body_size", TYPE["body"].size)
        while True:
            fits = all(self._fixed_height(b, w, size) <= (y1 - y0) for _, w, b in plans)
            if fits or size <= MIN_BODY_SIZE:
                break
            size -= 1
        if size < s.get("body_size", TYPE["body"].size):
            self.warn(f"body text auto-shrunk to {size}pt to fit")
        for x, w, blocks in plans:
            self._render_column(slide, x, w, y0, y1, blocks, size)
        if has_tracker:
            self.tracker(slide, s.get("section"))
        self.page_number(slide)
        self.notes(slide, s.get("notes"))

    def risks(self, s):
        g = GEO
        slide = self.new_slide("content")
        self.set_title(slide, s.get("title", "Risks & Mitigants"))
        rows = s["rows"]
        has_tracker = bool(self.sections) and s.get("section") is not None
        y0, y1 = 0.95, (g.body_y1_tracker if has_tracker else g.body_y1_plain)
        pitch = (y1 - y0) / len(rows)
        box_h = min(0.912, pitch - 0.42)
        for i, r in enumerate(rows):
            y = y0 + i * pitch
            yb = y + 0.323
            self.label_rect(slide, 0.083, yb, 3.713, box_h, r["risk"], TYPE["risk_label"], PALETTE["white"],
                            PALETTE["primary"], 2.25, shape=MSO_SHAPE.PENTAGON, name=f"Risk {i+1}", inset_tb=GEO.inset)
            self.label_rect(slide, 3.982, y, 2.254, 0.225, "Level of Risk", TYPE["mini_header"], PALETTE["ink"], PALETTE["ink"])
            self.label_rect(slide, 3.982, yb, 2.254, box_h, r.get("level", ""), TYPE["risk_level"],
                            PALETTE["primary"], PALETTE["panel"])
            self.label_rect(slide, 6.666, y, 6.224, 0.225, "Mitigation:", TYPE["mini_header"], PALETTE["ink"], PALETTE["ink"])
            mt = r.get("mitigation", "")
            msize = TYPE["risk_body"].size
            while msize > 10 and est_lines(mt, 6.0, msize) * line_h(msize) > box_h - 0.1:
                msize -= 0.5
            self.label_rect(slide, 6.666, yb, 6.224, box_h, mt, TYPE["risk_body"], PALETTE["panel"],
                            PALETTE["panel"], size=msize)
            c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(6.236), I(yb + box_h / 2), I(6.666), I(yb + box_h / 2))
            c.line.color.rgb = rgb(PALETTE["panel"])
            c.line.width = Pt(3)
            if i < len(rows) - 1:
                sep_y = yb + box_h + (pitch - 0.323 - box_h) / 2
                c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(3.929), I(sep_y), I(g.body_x1), I(sep_y))
                c.line.color.rgb = rgb(PALETTE["ink"])
                c.line.width = Pt(0.75)
        if has_tracker:
            self.tracker(slide, s.get("section"))
        self.page_number(slide)
        self.notes(slide, s.get("notes"))

    # ======================================================== block engine
    FLEX = ("image", "chart", "table", "logos")

    def _block_kind(self, b):
        for k in ("header", "caption", "bullets", "text", "image", "logos", "chart", "table",
                  "callouts", "metrics", "numbered", "spacer"):
            if k in b:
                return k
        raise ValueError(f"unknown block: {b}")

    def _bullet_items(self, items):
        """Flatten [str | {text, sub:[...]}] into [(level, text)]."""
        out = []
        for it in items:
            if isinstance(it, dict) and "text" not in it and len(it) == 1:
                # unquoted YAML like "- Danaos (NYSE: DAC) is ..." parses as {key: value}
                k, v = next(iter(it.items()))
                it = f"{k}: {v}"
            if isinstance(it, dict):
                out.append((0, it["text"]))
                out += [(1, t) for t in it.get("sub", [])]
            else:
                out.append((0, it))
        return out

    def _bullets_height(self, items, w, size, spacing=1.0, space_before=0):
        h = 2 * GEO.inset
        for lvl, t in self._bullet_items(items):
            text_w = w - 2 * GEO.inset - (GEO.bullet_marL if lvl == 0 else GEO.bullet_l2_marL)
            h += est_lines(t, text_w, size) * line_h(size, spacing) + space_before / 72
        return h

    def _height(self, b, w, size):
        """Natural height of a block, or None for flexible blocks."""
        k = self._block_kind(b)
        if k in ("header", "caption"):
            return GEO.header_h
        if k == "bullets":
            return self._bullets_height(b["bullets"], w, size, b.get("line_spacing", 1.0), b.get("space_before", 0))
        if k == "text":
            t = b["text"] if isinstance(b["text"], dict) else {"content": b["text"]}
            sz = t.get("size", size)
            return 2 * GEO.inset + est_lines(t["content"], w - 0.2, sz, t.get("small_caps", False)) * line_h(sz, t.get("line_spacing", 1.0))
        if k == "callouts":
            cols = b.get("cols", 2)
            rows = math.ceil(len(b["callouts"]) / cols)
            return rows * 0.706 + (rows - 1) * 0.1
        if k == "metrics":
            return b.get("height", 1.0)
        if k == "numbered":
            ts = TYPE["summary_body"]
            h = 0
            for it in b["numbered"]:
                h += 0.379 + 0.06 + 2 * 0.05 + est_lines(it.get("text", ""), w - 0.2, ts.size) * line_h(ts.size) + 0.12
            return h
        if k == "spacer":
            return float(b["spacer"])
        return b.get("height")  # flexible unless pinned

    def _fixed_height(self, blocks, w, size):
        total = 0.0
        for i, b in enumerate(blocks):
            h = self._height(b, w, size)
            total += h if h is not None else 1.0  # reserve 1in minimum for flex blocks
            if i < len(blocks) - 1:
                total += GEO.header_gap if self._block_kind(b) == "header" else GEO.block_gap
        return total

    def _render_column(self, slide, x, w, y0, y1, blocks, size):
        heights = [self._height(b, w, size) for b in blocks]
        gaps = [GEO.header_gap if self._block_kind(b) == "header" else GEO.block_gap for b in blocks[:-1]]
        fixed = sum(h for h in heights if h is not None) + sum(gaps)
        flex = [i for i, h in enumerate(heights) if h is None]
        avail = y1 - y0 - fixed
        if flex:
            wts = [blocks[i].get("weight", 1) for i in flex]
            for i, wt in zip(flex, wts):
                heights[i] = max(avail * wt / sum(wts), 0.6)
            if avail < 0.6 * len(flex):
                self.warn("not enough room for image/chart/table blocks; trim text or pin heights")
        elif avail < -0.05:
            self.warn(f"column content is ~{-avail:.2f} in taller than the body area")
        y = y0
        for i, (b, h) in enumerate(zip(blocks, heights)):
            used = getattr(self, "_b_" + self._block_kind(b))(slide, x, y, w, h, b, size)
            y += (used if used else h) + (gaps[i] if i < len(gaps) else 0)

    # ---- individual blocks
    def _b_header(self, slide, x, y, w, h, b, size):
        ts = TYPE["header_bar"]
        sz = b.get("size", ts.size)
        self.label_rect(slide, x, y, w, h, b["header"], ts, PALETTE["primary"], PALETTE["active"], 0.75,
                        name=f"Header: {plain(b['header'])}", size=sz)

    def _b_caption(self, slide, x, y, w, h, b, size):
        self.label_rect(slide, x, y, w, h, b["caption"], TYPE["header_bar"], PALETTE["neutral"],
                        name=f"Caption: {plain(b['caption'])}")

    def _b_spacer(self, *a):
        pass

    def _b_bullets(self, slide, x, y, w, h, b, size):
        tb = self.textbox(slide, x, y, w, h, "Bullets")
        tf = tb.text_frame
        ts = replace(TYPE["body"], color=b.get("color", TYPE["body"].color))
        for i, (lvl, t) in enumerate(self._bullet_items(b["bullets"])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if b.get("line_spacing"):
                p.line_spacing = b["line_spacing"]
            if b.get("space_before"):
                p.space_before = Pt(b["space_before"])
            set_bullet(p, lvl, ts.color, size, self.font)
            add_runs(p, t, ts, self.font, size=size)

    def _b_text(self, slide, x, y, w, h, b, size):
        t = b["text"] if isinstance(b["text"], dict) else {"content": b["text"]}
        ts = TypeStyle(t.get("size", size), t.get("bold", False), t.get("color", TYPE["body"].color),
                       small_caps=t.get("small_caps", False))
        tb = self.textbox(slide, x, y, w, h, "Text", anchor=MSO_ANCHOR.MIDDLE if t.get("middle") else MSO_ANCHOR.TOP)
        for i, para in enumerate(str(t["content"]).split("\n")):
            p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph()
            p.alignment = ALIGN[t.get("align", "l")]
            if t.get("line_spacing"):
                p.line_spacing = t["line_spacing"]
            no_bullet(p)
            add_runs(p, para, ts, self.font)

    def _fit_image(self, path, w, h):
        with Image.open(path) as im:
            ar = im.width / im.height
        iw, ih = (w, w / ar) if w / h < ar else (h * ar, h)
        return iw, ih

    def _b_image(self, slide, x, y, w, h, b, size):
        path = self._path(b["image"])
        if not path or not os.path.exists(path):
            self.warn(f"image not found: {b['image']}")
            self.label_rect(slide, x, y, w, h, f"Missing image: {b['image']}", TYPE["callout"], PALETTE["neutral"])
            return
        iw, ih = self._fit_image(path, w, h)
        px = x + (w - iw) / 2
        pic = slide.shapes.add_picture(path, I(px), I(y), I(iw), I(ih))
        if b.get("border"):
            pic.line.color.rgb = rgb(PALETTE["primary"])
            pic.line.width = Pt(1.5)
        if b.get("alt"):
            pic._element.nvPicPr.cNvPr.set("descr", b["alt"])
        return ih  # following blocks (e.g. a caption) sit directly under the picture

    def _b_logos(self, slide, x, y, w, h, b, size):
        paths = b["logos"]
        cols = b.get("cols", 2)
        rows = math.ceil(len(paths) / cols)
        cw, ch = w / cols, h / rows
        pad = b.get("padding", 0.15)
        for i, p in enumerate(paths):
            r, c = divmod(i, cols)
            full = self._path(p)
            if not os.path.exists(full):
                self.warn(f"logo not found: {p}")
                continue
            iw, ih = self._fit_image(full, cw - 2 * pad, ch - 2 * pad)
            slide.shapes.add_picture(full, I(x + c * cw + (cw - iw) / 2), I(y + r * ch + (ch - ih) / 2), I(iw), I(ih))

    def _b_callouts(self, slide, x, y, w, h, b, size):
        cols = b.get("cols", 2)
        gap = 0.25
        cw = (w - gap * (cols - 1)) / cols
        for i, it in enumerate(b["callouts"]):
            r, c = divmod(i, cols)
            text = it if isinstance(it, str) else f"__{it['lead']}__ {it.get('text', '')}".strip()
            self.label_rect(slide, x + c * (cw + gap), y + r * 0.806, cw, 0.706, text, TYPE["callout"],
                            PALETTE["active"], inset_tb=GEO.inset, name="Stat Callout")

    def _b_metrics(self, slide, x, y, w, h, b, size):
        items = b["metrics"]
        gap = 0.12
        cw = (w - gap * (len(items) - 1)) / len(items)
        for i, it in enumerate(items):
            text = it if isinstance(it, str) else (f"{it['label']}: {it['value']}" if it.get("label") else it["value"])
            ts = TYPE["metric"]
            sz = ts.size
            while sz > 16 and est_lines(text, cw - 0.2, sz, True) * line_h(sz) > h - 0.1:
                sz -= 1
            self.label_rect(slide, x + i * (cw + gap), y, cw, h, text, ts, PALETTE["panel"],
                            PALETTE["panel_line"], 1.0, size=sz, name="Metric Tile")

    def _b_numbered(self, slide, x, y, w, h, b, size):
        ts = TYPE["summary_body"]
        for i, it in enumerate(b["numbered"]):
            self.label_rect(slide, x, y, 0.635, 0.379, str(i + 1), TypeStyle(18, False, PALETTE["white"]),
                            PALETTE["ink"], PALETTE["ink"])
            self.label_rect(slide, x + 0.651, y, w - 0.651, 0.379, it["title"], TYPE["numbered_title"],
                            PALETTE["panel"], PALETTE["panel"])
            th = 2 * 0.05 + est_lines(it.get("text", ""), w - 0.2, ts.size) * line_h(ts.size)
            tb = self.textbox(slide, x, y + 0.379 + 0.06, w, th, inset_tb=0.05)
            p = tb.text_frame.paragraphs[0]
            no_bullet(p)
            add_runs(p, it.get("text", ""), ts, self.font)
            y += 0.379 + 0.06 + th + 0.12

    # ---- charts
    CHART_TYPES = {
        "bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "stacked": XL_CHART_TYPE.COLUMN_STACKED, "barh": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
    }

    def _chart_font(self, font_obj, size=None, bold=False, color=None):
        ts = TYPE["chart"]
        font_obj.name = self.font
        font_obj.size = Pt(size or ts.size)
        font_obj.bold = bold
        font_obj.color.rgb = rgb(color or ts.color)

    def _b_chart(self, slide, x, y, w, h, b, size):
        c = b["chart"]
        kind = c.get("type", "bar")
        cd = CategoryChartData(number_format=c.get("number_format", "General"))
        cd.categories = c["categories"]
        for s in c["series"]:
            cd.add_series(s["name"], s["values"])
        src_h = 0.25 if c.get("source") else 0.0
        gf = slide.shapes.add_chart(self.CHART_TYPES[kind], I(x), I(y), I(w), I(h - src_h), cd)
        ch = gf.chart
        self._chart_font(ch.font)
        if c.get("title"):
            ch.has_title = True
            tf = ch.chart_title.text_frame
            tf.text = c["title"]
            self._chart_font(tf.paragraphs[0].runs[0].font, size=14, bold=True)
        else:
            ch.has_title = False
        pie = kind in ("pie", "doughnut")
        multi = len(c["series"]) > 1
        ch.has_legend = c.get("legend", pie or multi)
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            self._chart_font(ch.legend.font, size=11)
        plot = ch.plots[0]
        colors = c.get("colors", CHART_SERIES)
        if pie:
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
                pt.format.line.color.rgb = rgb(PALETTE["white"])
        else:
            if hasattr(plot, "gap_width"):
                plot.gap_width = c.get("gap_width", 60)
            if kind == "stacked":
                plot.overlap = 100
            for i, s in enumerate(plot.series):
                col = rgb(c["series"][i].get("color", colors[i % len(colors)]))
                if kind == "line":
                    s.format.line.color.rgb = col
                    s.format.line.width = Pt(2.25)
                    s.smooth = False
                    s.marker.style = XL_MARKER_STYLE.NONE
                else:
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = col
            va, ca = ch.value_axis, ch.category_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb(PALETTE["grid"])
            va.major_gridlines.format.line.width = Pt(0.5)
            va.format.line.fill.background()
            va.visible = c.get("value_axis", True)
            ca.format.line.color.rgb = rgb(PALETTE["grid"])
            ca.has_major_gridlines = False
            for ax in (va, ca):
                self._chart_font(ax.tick_labels.font, size=11)
            if c.get("value_format"):
                va.tick_labels.number_format = c["value_format"]
                va.tick_labels.number_format_is_linked = False
        if c.get("data_labels", True):
            plot.has_data_labels = True
            dl = plot.data_labels
            inside = pie or kind == "stacked"
            self._chart_font(dl.font, size=11, color=PALETTE["title"])
            dl.number_format = c.get("label_format", c.get("number_format", "General"))
            dl.number_format_is_linked = False
            if pie:
                dl.position = XL_LABEL_POSITION.CENTER
            elif kind == "line":
                dl.position = XL_LABEL_POSITION.ABOVE
            elif kind == "stacked":
                dl.position = XL_LABEL_POSITION.CENTER
            else:
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            if inside:  # labels sit on the fill: pick white or gray per slice/segment
                fmt = dl.number_format
                if pie:
                    for i, pt in enumerate(plot.series[0].points):
                        self._chart_font(pt.data_label.font, size=11, color=ink_on(colors[i % len(colors)]))
                        pt.data_label.position = XL_LABEL_POSITION.CENTER
                else:
                    for i, sr in enumerate(plot.series):
                        col = c["series"][i].get("color", colors[i % len(colors)])
                        sdl = sr.data_labels
                        sdl.show_value = True
                        sdl.number_format = fmt
                        sdl.number_format_is_linked = False
                        sdl.position = XL_LABEL_POSITION.CENTER
                        self._chart_font(sdl.font, size=11, color=ink_on(col))
        if c.get("source"):
            tb = self.textbox(slide, x, y + h - src_h, w, src_h, inset_tb=0.0)
            p = tb.text_frame.paragraphs[0]
            add_runs(p, "Source: " + c["source"], TypeStyle(9, False, PALETTE["page_number"]), self.font)

    # ---- tables
    def _cell_border(self, cell, side, color, pt):
        tcPr = cell._tc.get_or_add_tcPr()
        tag = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}[side]
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.Element(qn(tag))
        if color:
            ln.set("w", str(int(Pt(pt))))
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        else:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
        idx = 0
        for t in order[: order.index(tag)]:
            if tcPr.find(qn(t)) is not None:
                idx += 1
        tcPr.insert(idx, ln)

    def _b_table(self, slide, x, y, w, h, b, size):
        t = b["table"]
        header = t.get("columns")
        rows = [r if isinstance(r, dict) else {"cells": r} for r in t["rows"]]
        n_rows = len(rows) + (1 if header else 0)
        n_cols = len(header or rows[0]["cells"])
        fsize = t.get("size", TYPE["table"].size)
        row_h = lambda s: s * 1.9 / 72  # noqa: E731
        while fsize > MIN_TABLE_SIZE and n_rows * row_h(fsize) > h:
            fsize -= 0.5
        if n_rows * row_h(fsize) > h + 0.05:
            self.warn(f"table needs ~{n_rows * row_h(fsize):.2f} in but has {h:.2f} in")
        th = min(h, n_rows * row_h(fsize))
        gf = slide.shapes.add_table(n_rows, n_cols, I(x), I(y), I(w), I(th))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for a in ("firstRow", "bandRow"):
            tblPr.set(a, "0")
        sid = tblPr.find(qn("a:tableStyleId"))
        if sid is None:
            sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
        sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
        widths = t.get("col_widths")
        if widths:
            tot = sum(widths)
            for i, cw in enumerate(widths):
                tbl.columns[i].width = I(w * cw / tot)
        for r in tbl.rows:
            r.height = I(th / n_rows)
        align = t.get("align") or (["l"] + ["r"] * (n_cols - 1))
        all_rows = ([{"cells": header, "style": "header"}] if header else []) + rows
        for ri, row in enumerate(all_rows):
            style = row.get("style", "")
            for ci in range(n_cols):
                cell = tbl.cell(ri, ci)
                val = row["cells"][ci] if ci < len(row["cells"]) else ""
                cell.margin_left = cell.margin_right = I(0.06)
                cell.margin_top = cell.margin_bottom = I(0.02)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                p = tf.paragraphs[0]
                p.alignment = ALIGN[align[ci]]
                color = PALETTE["white"] if style == "header" else TYPE["table"].color
                bold = style in ("header", "total", "highlight")
                indent = style == "indent" and ci == 0
                ts = TypeStyle(fsize, bold, color)
                add_runs(p, ("    " if indent else "") + str(val), ts, self.font)
                if style == "italic":
                    for r in p.runs:
                        r.font.italic = True
                fill = {"header": PALETTE["primary"], "highlight": PALETTE["highlight"]}.get(style)
                if fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(fill)
                else:
                    cell.fill.background()
                for side in "LR":
                    self._cell_border(cell, side, None, 0)
                top = PALETTE["primary"] if style == "total" else None
                self._cell_border(cell, "T", top, 1.0)
                self._cell_border(cell, "B", PALETTE["grid"] if style != "header" else None, 0.5)

    # ======================================================== driver
    def build(self, out_path):
        dispatch = {"cover": self.cover, "agenda": self.agenda, "divider": self.divider,
                    "content": self.content, "risks": self.risks}
        for s in self.spec["slides"]:
            kind = s.get("type", "content")
            if kind not in dispatch:
                raise ValueError(f"unknown slide type '{kind}'")
            dispatch[kind](s)
        self.prs.core_properties.title = self.spec.get("title", self.brand.target_name)
        self.prs.core_properties.author = self.spec.get("author", "")
        os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
        self.prs.save(out_path)
        return self.warnings


def build_deck(spec: dict, base_dir: str, out_path: str) -> list[str]:
    return DeckBuilder(spec, base_dir).build(out_path)
