#!/usr/bin/env python3
"""Check a built deck for what PowerPoint refuses and for anything not editable.

    python3 tools/check_pptx.py ../DKS_Pitch_Deck.pptx [--native-only]

Schema validators miss these, and LibreOffice renders them without complaint:
* data-label positions the chart type does not offer (PowerPoint refuses the
  file, or reports "can't read" for a downloaded copy)
* chart axis ids outside 0..2147483647
* duplicate shape ids, table rows that disagree with the table grid
* external links, macros, ActiveX or OLE objects
It also checks editability: no pictures or groups (FAIL with --native-only,
WARN otherwise), every slide has a real title, every chart has alt text and an
embedded workbook that matches the chart, and the document properties are the
deck's own rather than python-pptx's template values.
"""
from __future__ import annotations

import argparse
import io
import re
import sys
import zipfile
from collections import Counter

import openpyxl
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
# Label positions PowerPoint's Format Data Labels pane offers, by chart type.
LEGAL = {"barChart": {"clustered": {"ctr", "inBase", "inEnd", "outEnd"},
                      "stacked": {"ctr", "inBase", "inEnd"}, "percentStacked": {"ctr", "inBase", "inEnd"}},
         "lineChart": {"*": {"ctr", "l", "r", "t", "b"}},
         "pieChart": {"*": {"bestFit", "ctr", "inEnd", "outEnd"}},
         "doughnutChart": {"*": set()}, "areaChart": {"*": set()}, "radarChart": {"*": set()}}


def package_checks(z: zipfile.ZipFile, fail):
    for name in sorted(n for n in z.namelist() if re.fullmatch(r"ppt/charts/chart\d+\.xml", n)):
        t = etree.fromstring(z.read(name))
        for grp in t.find(".//c:plotArea", NS):
            kind = etree.QName(grp).localname
            if kind not in LEGAL:
                continue
            g = grp.find("c:grouping", NS)
            allowed = LEGAL[kind].get(g.get("val") if g is not None else "*", LEGAL[kind].get("*"))
            bad = Counter(p.get("val") for p in grp.iter(f"{{{NS['c']}}}dLblPos") if p.get("val") not in allowed)
            for pos, n in bad.items():
                fail(f"{name}: {n} data label(s) at '{pos}', which a {kind} does not offer")
        for ax in t.iter(f"{{{NS['c']}}}axId", f"{{{NS['c']}}}crossAx"):
            if not 0 <= int(ax.get("val")) <= 2147483647:
                fail(f"{name}: axis id {ax.get('val')} outside 0..2147483647")
                break
    for name in z.namelist():
        if re.fullmatch(r"ppt/(slides|slideLayouts|slideMasters)/\w+\.xml", name):
            t = etree.fromstring(z.read(name))
            ids = Counter(e.get("id") for e in t.iter(f"{{{NS['p']}}}cNvPr"))
            if any(n > 1 for n in ids.values()):
                fail(f"{name}: duplicate shape ids {[i for i, n in ids.items() if n > 1]}")
            for tbl in t.iter(f"{{{NS['a']}}}tbl"):
                ncol = len(tbl.findall("a:tblGrid/a:gridCol", NS))
                for r, tr in enumerate(tbl.findall("a:tr", NS)):
                    if len(tr.findall("a:tc", NS)) != ncol:
                        fail(f"{name}: table row {r} has {len(tr.findall('a:tc', NS))} cells for {ncol} columns")
        if name.endswith(".rels") and b'TargetMode="External"' in z.read(name):
            fail(f"{name}: external link (recipients get an update-links prompt)")
    ct = z.read("[Content_Types].xml").decode()
    for bad in ("macroEnabled", "vbaProject", "activeX", "oleObject"):
        if bad in ct:
            fail(f"package contains {bad} content")


def slide_checks(prs, fail, warn, native_only):
    kinds = Counter()
    for i, s in enumerate(prs.slides, 1):
        title = s.shapes.title
        if title is None or not title.text_frame.text.strip():
            fail(f"slide {i}: no text in a title placeholder (Outline view and screen readers need it)")
        for sh in s.shapes:
            st = sh.shape_type
            if st in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.GROUP):
                (fail if native_only else warn)(f"slide {i}: {st} '{sh.name}' is not editable as text or data")
            if sh.is_placeholder:
                kinds["title placeholder"] += 1
            elif sh.has_chart:
                kinds["chart"] += 1
                chart_checks(i, sh, fail)
            elif sh.has_table:
                kinds["table"] += 1
            elif st == MSO_SHAPE_TYPE.TEXT_BOX:
                kinds["text box"] += 1
            elif st == MSO_SHAPE_TYPE.AUTO_SHAPE:
                kinds["shape"] += 1
            else:
                kinds[str(st).split(" ")[0].lower()] += 1
    return kinds


def chart_checks(i, sh, fail):
    if not sh._element.nvGraphicFramePr.cNvPr.get("descr"):
        fail(f"slide {i}: chart '{sh.name}' has no alt text")
    ch = sh.chart
    ws = openpyxl.load_workbook(io.BytesIO(ch.part.chart_workbook.xlsx_part.blob)).worksheets[0]
    cats = list(ch.plots[0].categories)
    as_num = lambda vs: [None if v is None else float(v) for v in vs]  # noqa: E731
    if [ws.cell(row=r + 2, column=1).value for r in range(len(cats))] != cats:
        fail(f"slide {i}: chart '{sh.name}' categories differ from its workbook (Edit Data would change the chart)")
    for j, ser in enumerate(ch.plots[0].series):
        col = [ws.cell(row=r + 2, column=j + 2).value for r in range(len(cats))]
        if as_num(ser.values) != as_num(col) or ws.cell(row=1, column=j + 2).value != ser.name:
            fail(f"slide {i}: chart '{sh.name}' series {j + 1} differs from its workbook")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--native-only", action="store_true", help="fail on pictures and groups")
    args = ap.parse_args()
    fails, warns = [], []
    package_checks(zipfile.ZipFile(args.deck), fails.append)
    prs = Presentation(args.deck)
    kinds = slide_checks(prs, fails.append, warns.append, args.native_only)
    cp = prs.core_properties
    if cp.last_modified_by == "Steve Canny" or "python-pptx" in (cp.comments or "") or \
            (cp.created and cp.created.year < 2020):
        fails.append(f"document properties are python-pptx's template values ({cp.last_modified_by}, {cp.created})")
    print(f"{args.deck}: {len(prs.slides)} slides | " + ", ".join(f"{n} {k}" for k, n in sorted(kinds.items())))
    print(f"properties: title '{cp.title}', author '{cp.author}', created {cp.created}")
    for w in warns:
        print("  WARN", w)
    for f in fails:
        print("  FAIL", f)
    print("PASS" if not fails else f"FAIL ({len(fails)})")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
