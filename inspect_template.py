"""Reconnaissance pass on the YIS template — writes template_structure.json."""
from pptx import Presentation
import json, os

SRC = "Chen_Liu_Rochester_Adams_Troy_High_MI_PRESENTATION__1___2_.pptx"
prs = Presentation(SRC)
print(f"Slides: {len(prs.slides)}")
print(f"Size: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} in")

structure = []
for i, slide in enumerate(prs.slides, 1):
    sd = {"slide_num": i, "shapes": []}
    for shape in slide.shapes:
        info = {
            "shape_id": shape.shape_id,
            "name": shape.name,
            "type": str(shape.shape_type),
            "left_in": round(shape.left/914400, 2) if shape.left is not None else None,
            "top_in": round(shape.top/914400, 2) if shape.top is not None else None,
            "w_in": round(shape.width/914400, 2) if shape.width is not None else None,
            "h_in": round(shape.height/914400, 2) if shape.height is not None else None,
            "has_text": shape.has_text_frame,
        }
        if shape.has_text_frame:
            t = shape.text_frame.text
            info["text"] = t
            info["n_paragraphs"] = len(shape.text_frame.paragraphs)
        if shape.has_table:
            tbl = shape.table
            info["table"] = {
                "rows": len(tbl.rows), "cols": len(tbl.columns),
                "cells": [[c.text for c in row.cells] for row in tbl.rows]
            }
        sd["shapes"].append(info)
    structure.append(sd)

with open("template_structure.json", "w") as f:
    json.dump(structure, f, indent=2, default=str)
print("Wrote template_structure.json", os.path.getsize("template_structure.json"), "bytes")
